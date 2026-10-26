import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.finance.services import payments as pay_service
from apps.finance.tests.factories import AssessmentFactory, DuesTypeFactory, pay
from apps.reporting import services
from apps.reporting.exports import _safe

pytestmark = pytest.mark.django_db


@pytest.fixture
def ledger(department, session, membership, make_membership, make_user):
    finsec = make_user("financial_secretary")
    dues = DuesTypeFactory(department=department, academic_session=session, created_by=finsec)
    a1 = AssessmentFactory(dues_type=dues, department_membership=membership)
    a2 = AssessmentFactory(dues_type=dues, department_membership=make_membership(level_code="Level 300"))
    pay(a1, "150.00")
    pay(a2, "50.00", method="mobile_money")
    pay(a2, "100.00", status="pending")
    return dues


def test_session_totals_and_levels(department, session, ledger):
    totals = services.session_totals(department, session)
    assert (totals.assessed, totals.collected, totals.outstanding) == (
        Decimal("300"),
        Decimal("200"),
        Decimal("100"),
    )
    assert totals.members == 2
    assert round(totals.collection_rate, 2) == Decimal("66.67")
    levels = {row["level"]: row["rate"] for row in services.by_level(department, session)}
    assert levels == {"Level 200": Decimal("100"), "Level 300": Decimal("50") / Decimal("150") * 100}


def test_income_excludes_pending_payments(department, ledger):
    rows = services.payments_by_method(department)
    assert sum(r["total"] for r in rows) == Decimal("200.00")
    assert {r["method_label"] for r in rows} == {"Cash", "Mobile money"}


def test_monthly_income_is_zero_filled_and_ordered(department, ledger):
    series = services.monthly_income(department, months=12)
    assert len(series) == 12
    assert series[-1]["month"] == datetime.date.today().replace(day=1)
    assert series[-1]["total"] == Decimal("200.00")
    assert series[0]["total"] == 0


def test_dashboard_by_role(member, make_user, login, ledger):
    student = login(member).get("/")
    assert "member" in student.context and "finance" not in student.context
    treasurer = login(make_user("treasurer")).get("/")
    assert treasurer.context["finance"]["totals"].collected == Decimal("200.00")
    assert treasurer.context["finance"]["queues"]["pending_payments"] == 1


def test_report_permissions(member, make_user, login, ledger):
    assert login(member).get("/reports/finance/").status_code == 403
    assert login(make_user("auditor")).get("/reports/finance/").status_code == 200
    finsec = login(make_user("financial_secretary"))
    assert finsec.get("/reports/finance/export/balances/").status_code == 403  # no export permission


def test_exports_are_audited(make_user, login, ledger, session):
    client = login(make_user("treasurer"))
    csv_response = client.get("/reports/finance/export/balances/", {"session": session.pk, "format": "csv"})
    assert csv_response.status_code == 200
    lines = csv_response.content.decode("utf-8-sig").strip().splitlines()
    assert lines[0].startswith("Student number,Name,Email,Level")
    assert len(lines) == 3

    xlsx = client.get("/reports/finance/export/payments/", {"format": "xlsx"})
    sheet = load_workbook(__import__("io").BytesIO(xlsx.content)).active
    assert sheet.max_row == 4  # header + 3 payments (pending included, labelled)
    assert AuditEvent.objects.filter(action_key="report.exported").count() == 2


def test_formula_cells_are_neutralised():
    assert _safe("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert _safe("+233") == "'+233"
    assert _safe("Ama") == "Ama"
    assert _safe(Decimal("5")) == Decimal("5")


def test_anonymous_home_shows_sign_in(client, db):
    response = client.get("/")
    assert response.status_code == 200
    assert "Sign in" in response.content.decode()


def test_recorded_pending_payment_shows_in_queue(make_user, login, membership, ledger):
    finsec = make_user("financial_secretary")
    pay_service.record_offline_payment(membership, actor=finsec, method="cash", amount="5")
    page = login(make_user("treasurer")).get("/")
    assert page.context["finance"]["queues"]["pending_payments"] == 2
