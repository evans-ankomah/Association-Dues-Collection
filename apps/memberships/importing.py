"""Import a member roster (CSV or Excel) into the current academic session.

Columns (header names are case-insensitive): email, given_name, family_name, student_number,
phone_number, programme_name, level_code. Only email, given_name and family_name are required.

Two steps: `parse_roster` validates every row and reports what would happen, then
`apply_roster` commits the valid rows in one transaction.
"""

import csv
import io
from dataclasses import asdict, dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

from apps.access.models import Role, RoleAssignment
from apps.accounts.models import User, UserProfile
from apps.audit.services import record
from apps.common.phone import normalize_gh_phone

from .models import DepartmentMembership

COLUMNS = [
    "email",
    "given_name",
    "family_name",
    "student_number",
    "phone_number",
    "programme_name",
    "level_code",
]
REQUIRED = ["email", "given_name", "family_name"]
ALIASES = {
    "first_name": "given_name",
    "first name": "given_name",
    "surname": "family_name",
    "last_name": "family_name",
    "last name": "family_name",
    "index_number": "student_number",
    "student number": "student_number",
    "student_id": "student_number",
    "phone": "phone_number",
    "programme": "programme_name",
    "program": "programme_name",
    "level": "level_code",
}
MAX_ROWS = 2000


class Action:
    NEW_USER = "new_user"
    NEW_MEMBERSHIP = "new_membership"
    UPDATE = "update"
    ERROR = "error"


@dataclass
class RosterRow:
    line: int
    email: str = ""
    given_name: str = ""
    family_name: str = ""
    student_number: str = ""
    phone_number: str = ""
    programme_name: str = ""
    level_code: str = ""
    action: str = ""
    errors: list = field(default_factory=list)


def _normalise_header(name):
    key = (name or "").strip().lower()
    return ALIASES.get(key, key.replace(" ", "_"))


def _read_rows(upload):
    name = (getattr(upload, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = [_normalise_header(str(h or "")) for h in next(rows, [])]
        for values in rows:
            yield dict(zip(headers, ["" if v is None else str(v).strip() for v in values], strict=False))
        return
    raw = upload.read(5_000_000)
    text = raw.decode("utf-8-sig", errors="replace") if isinstance(raw, bytes) else raw
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [_normalise_header(h) for h in (reader.fieldnames or [])]
    for row in reader:
        yield {k: (v or "").strip() for k, v in row.items() if k}


def parse_roster(upload, *, department, session):
    rows, seen_emails, seen_numbers = [], set(), set()
    institution = department.institution
    for index, data in enumerate(_read_rows(upload), start=2):
        if index > MAX_ROWS + 1:
            raise ValueError(f"The file has more than {MAX_ROWS} rows. Split it and import in parts.")
        if not any(data.get(c) for c in COLUMNS):
            continue
        row = RosterRow(line=index, **{c: data.get(c, "") for c in COLUMNS})
        row.email = row.email.lower()
        for column in REQUIRED:
            if not getattr(row, column):
                row.errors.append(f"{column.replace('_', ' ')} is missing")
        if row.email:
            try:
                validate_email(row.email)
            except ValidationError:
                row.errors.append("email isn't valid")
            if row.email in seen_emails:
                row.errors.append("email appears twice in the file")
            seen_emails.add(row.email)
        if row.phone_number:
            try:
                row.phone_number = normalize_gh_phone(row.phone_number) or ""
            except ValueError:
                row.errors.append("phone number isn't a Ghanaian number")
        if row.student_number:
            if row.student_number in seen_numbers:
                row.errors.append("student number appears twice in the file")
            seen_numbers.add(row.student_number)
            clash = UserProfile.objects.filter(
                institution=institution, student_number=row.student_number
            ).exclude(user__email=row.email)
            if clash.exists():
                row.errors.append("student number belongs to a different email")

        if row.errors:
            row.action = Action.ERROR
        else:
            user = User.objects.filter(email=row.email).first()
            if user is None:
                row.action = Action.NEW_USER
            elif DepartmentMembership.objects.filter(
                user=user, department=department, academic_session=session
            ).exists():
                row.action = Action.UPDATE
            else:
                row.action = Action.NEW_MEMBERSHIP
        rows.append(row)
    if not rows:
        raise ValueError("The file has no member rows. Check it has an email column.")
    return rows


def summarise(rows):
    counts = {Action.NEW_USER: 0, Action.NEW_MEMBERSHIP: 0, Action.UPDATE: 0, Action.ERROR: 0}
    for row in rows:
        counts[row.action] += 1
    return counts


def to_session(rows):
    return [asdict(r) for r in rows]


def from_session(data):
    return [RosterRow(**r) for r in data]


@transaction.atomic
def apply_roster(rows, *, department, session, actor):
    """Create or update the valid rows. Returns the users created (for welcome emails)."""
    institution = department.institution
    member_role = Role.objects.get(role_code="student_member")
    now = timezone.now()
    created_users, counts = [], {"created": 0, "joined": 0, "updated": 0, "skipped": 0}
    for row in rows:
        if row.action == Action.ERROR:
            counts["skipped"] += 1
            continue
        user, user_created = User.objects.get_or_create(email=row.email, defaults={"is_active": True})
        if user_created:
            user.set_unusable_password()
            user.save(update_fields=["password"])
            created_users.append(user)
            counts["created"] += 1
        profile, _ = UserProfile.objects.get_or_create(
            user=user,
            defaults={
                "institution": institution,
                "given_name": row.given_name,
                "family_name": row.family_name,
            },
        )
        profile.given_name = row.given_name
        profile.family_name = row.family_name
        if row.student_number:
            profile.student_number = row.student_number
        if row.phone_number:
            profile.phone_number = row.phone_number
        profile.save()

        membership, membership_created = DepartmentMembership.objects.get_or_create(
            user=user,
            department=department,
            academic_session=session,
            defaults={
                "programme_name": row.programme_name or None,
                "level_code": row.level_code or None,
                "joined_on": timezone.localdate(),
                "verified_by": actor,
                "verified_at": now,
            },
        )
        if membership_created:
            counts["joined"] += 0 if user_created else 1
        else:
            membership.programme_name = row.programme_name or membership.programme_name
            membership.level_code = row.level_code or membership.level_code
            membership.verified_by = actor
            membership.verified_at = now
            membership.save()
            counts["updated"] += 1
        RoleAssignment.objects.get_or_create(
            user=user,
            role=member_role,
            department=department,
            academic_session=session,
            assignment_status=RoleAssignment.Status.ACTIVE,
            defaults={"starts_at": now, "assigned_by": actor},
        )
    record(
        "members.roster_imported",
        None,
        entity_type="department_memberships",
        actor=actor,
        department=department,
        after={**counts, "session": session.name},
    )
    return created_users, counts
